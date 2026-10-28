import io
import zipfile

from conftest import H, make_user


def _setup(client, admin, event):
    eid, h = event["id"], H(admin)
    base = f"/api/events/{eid}"
    sp = []
    for n, st in (("Dr. Ananya Rao", "Confirmed"), ("Rohit =Gupta", "Invited")):
        sp.append(client.post(f"{base}/speakers", headers=h, json={"full_name": n, "status": st, "bio": "Bio <b>& more</b>",
                                                                  "email": "x@example.com", "designation": "MD", "organisation": "Acme"}).json())
    day, hall = event["days"][0]["id"], event["halls"][0]["id"]
    for title, status in (("Public session", "published"), ("Secret draft", "draft")):
        r = client.post(f"{base}/sessions", headers=h, json={
            "day_id": day, "hall_id": hall, "title": title, "session_type": "Panel Discussion", "status": status,
            "start_time": "10:00", "end_time": "11:00", "writeup": "Writeup & <text>", "subtopics": ["Point one"],
            "internal_notes": "PRIVATE-NOTE", "people": [{"speaker_id": sp[0]["id"], "role": "moderator"}]})
        assert r.status_code == 201, r.text
    g = client.post(f"{base}/groups", headers=h, json={"name": "Organising Committee", "group_type": "OC"}).json()
    client.post(f"{base}/groups/{g['id']}/members", headers=h, json={"name": "Neha Singh", "role_in_group": "Chair"})
    g = client.post(f"{base}/groups/{g['id']}/meetings", headers=h, json={
        "title": "Kick-off", "meeting_date": "2027-01-10", "start_time": "15:00", "duration_minutes": 60}).json()
    return base, h, sp, g


def test_exports_produce_real_files(client, admin, event):
    base, h, sp, g = _setup(client, admin, event)
    gid, mid = g["id"], g["meetings"][0]["id"]
    for path, kind in [("agenda.pdf", b"%PDF"), ("runsheet.pdf", b"%PDF"), ("speakers.pdf", b"%PDF"),
                       (f"speakers/{sp[0]['id']}.pdf", b"%PDF"), (f"groups/{gid}/members.pdf", b"%PDF"),
                       (f"groups/{gid}/schedule.pdf", b"%PDF"), (f"groups/{gid}/meetings/{mid}.pdf", b"%PDF"),
                       ("agenda.docx", b"PK"), ("speakers.xlsx", b"PK"), (f"groups/{gid}/members.xlsx", b"PK")]:
        r = client.get(f"{base}/export/{path}", headers=h)
        assert r.status_code == 200, (path, r.text[:200])
        assert r.content.startswith(kind) and "attachment" in r.headers["content-disposition"], path


def test_agenda_options_control_content(client, admin, event):
    base, h, *_ = _setup(client, admin, event)

    def docx_text(q):
        r = client.get(f"{base}/export/agenda.docx{q}", headers=h)
        xml = zipfile.ZipFile(io.BytesIO(r.content)).read("word/document.xml").decode()
        return xml
    default = docx_text("")
    assert "Public session" in default and "Secret draft" not in default      # drafts off by default
    assert "PRIVATE-NOTE" not in default                                       # notes off by default
    assert "Writeup &amp; &lt;text&gt;" in default and "Point one" in default
    full = docx_text("?drafts=true&notes=true")
    assert "Secret draft" in full and "PRIVATE-NOTE" in full
    bare = docx_text("?writeups=false&subtopics=false")
    assert "Writeup" not in bare and "Point one" not in bare


def test_speaker_excel_filter_and_formula_safety(client, admin, event):
    from openpyxl import load_workbook
    base, h, *_ = _setup(client, admin, event)

    def rows(q):
        r = client.get(f"{base}/export/speakers.xlsx{q}", headers=h)
        return list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    all_rows, conf = rows(""), rows("?confirmed_only=true")
    assert len(all_rows) == 3 and len(conf) == 2 and conf[1][0] == "Dr. Ananya Rao"
    assert all_rows[0][5] == "Email" and "Rohit =Gupta" in [r[0] for r in all_rows]


def test_export_permissions_and_isolation(client, admin, event):
    base, h, sp, g = _setup(client, admin, event)
    outsider = make_user("outsider@example.com")
    assert client.get(f"{base}/export/agenda.pdf", headers=H(outsider)).status_code == 404   # other event: not visible
    assert client.get(f"{base}/export/agenda.pdf").status_code == 401
    assert client.get(f"{base}/export/speakers/99999.pdf", headers=h).status_code == 404
    assert client.get(f"{base}/export/groups/99999/members.pdf", headers=h).status_code == 404
    assert client.get(f"{base}/export/groups/{g['id']}/meetings/99999.pdf", headers=h).status_code == 404
    # a viewer may export (Section 2: "Viewer ... can export")
    inv = client.post(f"{base}/invitations", headers=h, json={"email": "view@example.com", "role": "viewer"})
    assert inv.status_code == 201
    token = inv.json()["link"].split("/invite/")[1]
    viewer = client.post("/api/auth/accept-invite", json={"token": token, "password": "password1", "full_name": "Vee"}).json()["access_token"]
    assert client.get(f"{base}/export/agenda.pdf", headers=H(viewer)).status_code == 200
    assert client.get(f"{base}/export/speakers.xlsx", headers=H(viewer)).status_code == 200
