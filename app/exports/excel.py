"""Excel exports: speakers (one row per speaker) and group members (one row per member)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import Group


def _sheet(title: str, headers: list[str], rows: list[list], widths: list[int]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    head_fill = PatternFill("solid", fgColor="0E1A42")  # IEML navy
    gold = Border(bottom=Side(style="medium", color="C19A45"))
    for c in ws[1]:
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), head_fill, gold
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(c.value, str):
                c.data_type = "s"  # text stays text: a value such as "=1+1" must never run as a formula
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def speakers_xlsx(speakers: list[dict]) -> bytes:
    heads = ["Full name", "Designation", "Organisation", "Status", "Remark", "Email", "Phone", "LinkedIn", "Website",
             "Bio", "Has photo", "Sessions", "Internal notes"]
    rows = []
    for s in speakers:
        sess = "\n".join(f"{r['role'].title()}: {r['title']} ({r['day']}, {r['start']}–{r['end']}, {r['hall']})" for r in s["sessions"])
        rows.append([s["name"], s["designation"], s["organisation"], s["status"], s["remark"], s["email"], s["phone"],
                     s["linkedin"], s["website"], s["bio"], "Yes" if s["has_photo"] else "No", sess, s["notes"]])
    return _sheet("Speakers", heads, rows, [26, 28, 28, 12, 28, 28, 16, 28, 28, 50, 10, 55, 36])


def members_xlsx(g: Group) -> bytes:
    heads = ["#", "Name", "Designation", "Organisation", "Role in group", "Email", "Phone", "Notes"]
    rows = [[i, m.name, m.designation, m.organisation, m.role_in_group, m.email, m.phone, m.notes] for i, m in enumerate(g.members, 1)]
    return _sheet(g.name, heads, rows, [5, 26, 28, 28, 22, 28, 16, 40])