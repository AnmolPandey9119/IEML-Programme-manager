"""Editable agenda (.docx), same content as the agenda PDF."""
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..models import Event
from .data import event_subtitle, event_title

INK, BRAND, MUTED, RED = RGBColor(0x10, 0x1C, 0x3D), RGBColor(0x1D, 0x3F, 0x91), RGBColor(0x5A, 0x64, 0x7F), RGBColor(0xA8, 0x39, 0x2B)
LOGO = Path(__file__).resolve().parent.parent / "assets" / "ieml-logo.png"


def _run(p, text, bold=False, size=10, color=INK, italic=False):
    r = p.add_run(text)
    r.bold, r.italic, r.font.size, r.font.color.rgb = bold, italic, Pt(size), color
    return r


def agenda_docx(ev: Event, days: list[dict], opts) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = sec.bottom_margin = Mm(18)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    try:  # company letterhead: IEML logo + name in the page header
        hp = sec.header.paragraphs[0]
        hp.add_run().add_picture(str(LOGO), height=Mm(11))
        _run(hp, "   India Exposition Mart Ltd.", True, 9, MUTED)
    except Exception:
        pass

    p = doc.add_paragraph()
    _run(p, event_title(ev), True, 24)
    if event_subtitle(ev):
        _run(doc.add_paragraph(), event_subtitle(ev), size=11, color=MUTED)
    p = doc.add_paragraph()
    _run(p, "Programme agenda", True, 13, BRAND)
    if opts.notes:
        _run(doc.add_paragraph(), "Contains internal notes. Do not share outside the team.", True, 9, RED)

    for d in days:
        h = doc.add_heading(level=1)
        _run(h, d["label"], True, 15, INK)
        h.paragraph_format.space_before = Pt(16)
        if not d["sessions"]:
            _run(doc.add_paragraph(), "No sessions scheduled for this day.", size=9, color=MUTED, italic=True)
            continue
        for s in d["sessions"]:
            meta = doc.add_paragraph()
            meta.paragraph_format.space_before, meta.paragraph_format.space_after = Pt(10), Pt(0)
            meta.paragraph_format.keep_with_next = True
            _run(meta, f"{s['start']} – {s['end']}", True, 9, BRAND)
            _run(meta, f"   {s['hall']}   ·   {s['type'].upper()}" + ("   ·   DRAFT" if s["draft"] else ""), False, 8.5, MUTED)
            t = doc.add_paragraph()
            t.paragraph_format.space_after, t.paragraph_format.keep_with_next = Pt(2), bool(s["writeup"] or s["people"])
            _run(t, s["title"], True, 12)
            if s["writeup"]:
                _run(doc.add_paragraph(), s["writeup"], size=10)
            for sub in s["subtopics"]:
                b = doc.add_paragraph(style="List Bullet")
                _run(b, sub, size=9.5)
            for pe in s["people"]:
                line = doc.add_paragraph()
                line.paragraph_format.space_after = Pt(0)
                _run(line, pe["role"].title() + "  ", True, 8.5, BRAND)
                _run(line, pe["name"], True, 9.5)
                sub = ", ".join(x for x in (pe["designation"], pe["organisation"]) if x)
                if sub:
                    _run(line, ", " + sub, size=9.5)
                if pe["status"] != "Confirmed":
                    _run(line, f" ({pe['status']})", size=9, color=RED)
            if s["notes"]:
                _run(doc.add_paragraph(), "Internal note: " + s["notes"], size=9, color=RED, italic=True)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, event_title(ev) + " · Programme agenda", size=8, color=MUTED)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()